"""Uji lapisan AI TANPA memanggil API sungguhan (transport tiruan). Jalankan: python tests/test_llm.py"""
import json, os, shutil, sys, tempfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import openpyxl
import data_processing as dp
from data_processing import DataFileError, build_database
from llm_service import LLMConfig, LLMError, LLMService, UsageGuard, load_config, parse_rewrite, validate_explanation
from search_service import merge_by_names, search_items

KEY = "KUNCI-RAHASIA-123"
CONN = build_database()
VOCAB = [r[0] for r in CONN.execute("SELECT nama FROM items")]


class Clock:
    def __init__(self): self.t = 1_000_000.0
    def __call__(self): return self.t


def reply(text, tokens=500):
    return {"candidates": [{"content": {"parts": [{"text": text}]}}], "usageMetadata": {"totalTokenCount": tokens}}


class Fake:
    """Transport tiruan: menghitung panggilan dan mengembalikan/men-raise sesuai skenario."""
    def __init__(self, *responses): self.responses, self.calls, self.last = list(responses), 0, None
    def __call__(self, url, headers, body, timeout):
        self.calls += 1; self.last = (url, headers, body)
        r = self.responses[min(self.calls - 1, len(self.responses) - 1)]
        if isinstance(r, Exception): raise r
        return r


def svc(*responses, **cfg):
    clock, fake = Clock(), Fake(*responses)
    return LLMService(LLMConfig(api_key=KEY, **cfg), transport=fake, clock=clock), fake, clock


# ---------------------------------------------------------------- pagar pemakaian
def test_rpm_disembunyikan_di_8_dari_15():
    s, f, clock = svc(reply("x"), min_interval_s=0)
    for _ in range(7):
        s.guard.reserve(); clock.t += 1
    assert s.available().allowed                       # 7 permintaan: masih boleh
    s.guard.reserve(); clock.t += 1
    st = s.available(); assert not st.allowed and st.reason == "rpm"   # 8 permintaan: disembunyikan
    clock.t += 61
    assert s.available().allowed                       # jendela 60 detik lewat: pulih


def test_rewrite_berhenti_lebih_awal_agar_tombol_tetap_ada():
    s, f, clock = svc(min_interval_s=0)
    for _ in range(6):
        s.guard.reserve(); clock.t += 1
    assert not s.available("rewrite").allowed and s.available("explain").allowed


def test_token_per_menit_dan_per_hari():
    s, f, clock = svc(min_interval_s=0)
    s.guard.reserve(); s.guard.record_tokens(150_000)
    st = s.available(); assert not st.allowed and st.reason == "tpm"
    s2, f2, c2 = svc(min_interval_s=0)
    for _ in range(400):
        s2.guard.reserve(); c2.t += 100
    assert s2.available().reason == "rpd"


def test_batas_sesi_dan_jeda_minimal():
    s, f, clock = svc()
    assert s.available(session_calls=8).reason == "session"
    s.guard.reserve()
    assert s.available().reason == "interval"
    clock.t += 4; assert s.available().allowed


def test_tanpa_kunci_atau_dimatikan_tidak_memanggil_api():
    fake = Fake(reply("x"))
    for cfg in (LLMConfig(api_key=""), LLMConfig(api_key=KEY, enabled=False)):
        o = LLMService(cfg, transport=fake).rewrite("bohlam", VOCAB)
        assert not o.ok
    assert fake.calls == 0


def test_cooldown_setelah_429_dan_circuit_breaker():
    s, f, clock = svc(LLMError("rate_limit", 120.0))
    o = s.rewrite("bohlam", VOCAB)
    assert not o.ok and o.reason == "rate_limit" and KEY not in o.message
    assert f.calls == 1
    clock.t += 10
    assert s.rewrite("bohlam", VOCAB).reason == "cooldown" and f.calls == 1   # tidak memanggil lagi
    s2, f2, c2 = svc(LLMError("network"), min_interval_s=0)
    for i in range(3):
        s2.rewrite(f"kata {i}", VOCAB); c2.t += 1
    assert s2.available().reason == "breaker"


# ---------------------------------------------------------------- keluaran LLM divalidasi
def test_rewrite_hanya_nama_yang_ada_pada_data():
    out = parse_rewrite('{"nama": ["Lampu LED 12 watt", "Rudal Balistik", "lampu led 12w"]}', VOCAB)
    assert out == ["Lampu LED 12 watt", "Lampu LED 12W"]
    assert parse_rewrite("bukan json", VOCAB) == [] and parse_rewrite('{"nama": "x"}', VOCAB) == []
    assert parse_rewrite('```json\n{"nama": ["Lakban bening"]}\n```', VOCAB) == ["Lakban bening"]


def test_rewrite_end_to_end_bohlam():
    s, f, _ = svc(reply('{"nama": ["Lampu LED 12 watt", "Lampu LED 12W"]}'))
    assert search_items(CONN, "bohlam").status == "not_found"          # aturan saja tidak menemukan
    o = s.rewrite("bohlam", VOCAB); assert o.ok and len(o.value) == 2
    merged = merge_by_names(CONN, "bohlam", o.value)
    assert {c.kode for c in merged.candidates} == {"BRG002", "BRG017"}
    assert "dipahami AI" in merged.candidates[0].reason


def test_rewrite_tidak_ada_padanan():
    s, f, _ = svc(reply('{"nama": []}'))
    o = s.rewrite("sepeda motor", VOCAB); assert o.ok and o.value == []


def test_penjelasan_valid_diterima_dan_di_cache():
    r = search_items(CONN, "lakban bening")
    s, f, _ = svc(reply("BRG019 tercatat 58 (Stok Barang, baris 20) dan BRG012 64 (Stok Barang, baris 13); angka berbeda, verifikasi fisik."))
    o = s.explain(r); assert o.ok and not o.cached
    o2 = s.explain(r); assert o2.ok and o2.cached and f.calls == 1     # panggilan kedua tidak memakai kuota
    assert KEY not in f.last[0] and f.last[1]["x-goog-api-key"] == KEY  # kunci hanya di header
    assert f.last[2]["generationConfig"]["maxOutputTokens"] == 600


def test_penjelasan_dengan_angka_karangan_ditolak():
    r = search_items(CONN, "lakban bening")
    s, f, _ = svc(reply("Stok lakban bening sebanyak 777 roll."))
    o = s.explain(r); assert not o.ok and o.reason == "validation" and "777" not in o.message


def test_penjelasan_gagal_kembali_ke_pesan_aman():
    r = search_items(CONN, "lakban bening")
    for err in (LLMError("rate_limit"), LLMError("network"), LLMError("http_500"), RuntimeError("boom " + KEY), {"candidates": []}):
        s, f, _ = svc(err)
        o = s.explain(r); assert not o.ok and KEY not in o.message and "tidak tersedia" in o.message


def test_penjelasan_hanya_untuk_hasil_ok():
    assert not LLMService(LLMConfig(api_key=KEY)).explain(search_items(CONN, "BRG099")).ok


def test_konfigurasi_dari_secrets():
    cfg = load_config({"GEMINI_API_KEY": " abc ", "LLM_ENABLED": "false", "LLM_RPM_HIDE_AT": "5"}.get)
    assert cfg.api_key == "abc" and not cfg.enabled and cfg.rpm_hide_at == 5 and cfg.rpm_limit == 15


def test_parser_toleran_bentuk_jawaban():
    from llm_service import parse_rewrite_ex
    ok = ["Lampu LED 12 watt", "Lampu LED 12W"]
    assert parse_rewrite_ex('Berikut hasilnya: {"nama": ["lampu led 12 WATT"]}', VOCAB)[0] == ok       # teks pembuka + huruf berbeda
    assert parse_rewrite_ex('["Lampu LED 12W"]', VOCAB)[0] == ok                                       # berupa list langsung
    assert parse_rewrite_ex('{"names": "Lampu LED 12W"}', VOCAB)[0] == ok                              # kunci/tipe berbeda
    assert parse_rewrite_ex("maaf saya tidak tahu", VOCAB) == ([], False)                              # bukan JSON
    assert parse_rewrite_ex('{"nama": ["Rudal"]}', VOCAB) == ([], True)                                # valid tetapi bukan nama pada data


def test_rewrite_no_match_dan_invalid_output_dibedakan():
    s, f, _ = svc(reply('{"nama": []}'))
    o = s.rewrite("sepeda motor", VOCAB); assert o.ok and o.value == [] and o.reason == "no_match" and "tidak menemukan" in o.message
    s2, f2, _ = svc(reply("saya tidak bisa membantu"))
    o2 = s2.rewrite("bohlam", VOCAB); assert not o2.ok and o2.reason == "invalid_output"


def test_debug_menyimpan_jawaban_mentah_tanpa_kunci():
    s, f, _ = svc(reply('{"nama": ["Lakban bening"]}', tokens=321))
    s.rewrite("selotip", VOCAB)
    d = s.last_debug; assert d["raw"].startswith('{"nama"') and d["tokens"] == 321 and d["parsed"] == ["Lakban bening"]
    assert KEY not in json.dumps(d)


# ---------------------------------------------------------------- pembaruan berkas Excel
def _copy_master(tmp):
    dst = os.path.join(tmp, "master.xlsx"); shutil.copy(dp.MASTER_FILE, dst); return dst


def test_pembaruan_excel_tambah_baris_kolom_dan_urutan():
    with tempfile.TemporaryDirectory() as tmp:
        path = _copy_master(tmp)
        wb = openpyxl.load_workbook(path); ws = wb["Stok Barang"]
        ws.append(["BRG021", "Gunting kertas", "Alat Tulis", "pcs", "Rak E-1", 30, 10, None])
        ws.append([None] * 8)                                  # baris kosong tidak boleh merusak
        ws.insert_cols(3); ws.cell(1, 3, "Kolom Tambahan Baru")  # kolom baru di tengah
        wb.save(path)
        items = dp.load_items(path)
        assert len(items) == 21 and "BRG021" in set(items["kode"])
        assert items[items["kode"] == "BRG021"].iloc[0]["sumber"] == "Stok Barang, baris 22"


def test_sheet_atau_kolom_wajib_hilang_memberi_pesan_jelas():
    with tempfile.TemporaryDirectory() as tmp:
        path = _copy_master(tmp)
        wb = openpyxl.load_workbook(path); wb["Catatan WA"].title = "WA"; wb.save(path)
        try: dp.load_wa_notes(path); assert False
        except DataFileError as e: assert "Catatan WA" in str(e) and "Sheet yang ada" in str(e)
        path2 = _copy_master(tmp)
        wb = openpyxl.load_workbook(path2); wb["Stok Barang"].cell(1, 1, "Kolom Lain"); wb.save(path2)
        try: dp.load_items(path2); assert False
        except DataFileError as e: assert "kode barang" in str(e)
    try: dp.load_end_stock(os.path.join(tmp, "tidak_ada.xlsx")); assert False
    except DataFileError as e: assert "tidak ditemukan" in str(e)


def test_nama_sheet_stok_akhir_bebas():
    with tempfile.TemporaryDirectory() as tmp:
        dst = os.path.join(tmp, "stok.xlsx"); shutil.copy(dp.ENDSTOCK_FILE, dst)
        wb = openpyxl.load_workbook(dst); wb.worksheets[0].title = "September 2026"; wb.save(dst)
        df = dp.load_end_stock(dst); assert len(df) == 8 and df.iloc[0]["sumber"].startswith("September 2026")


if __name__ == "__main__":
    fails = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            try: fn(); print("PASS", name)
            except Exception as e: fails += 1; print("FAIL", name, type(e).__name__, e)
    sys.exit(1 if fails else 0)
